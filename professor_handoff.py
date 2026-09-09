"""Build the professor-facing ROS_Spin report, PDF, summary, and ZIP package.

The script consumes only current paper-pipeline outputs and repository evidence.
Numerical result tables are never typed into the report by hand.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
DEFAULT_RESULTS = ROOT / "results" / "paper"
DEFAULT_DELIVERABLES = ROOT / "deliverables"
TITLE = "ROS Spin Professor Handoff"
SUBTITLE = "Corrected Doxorubicin Semiquinone Oxygen Spin Chemistry Model"
PROVENANCE_STATEMENT = (
    "All quantitative figures and tables in this document were generated directly "
    "from the ROS_Spin computational workflow or from values stored in its "
    "machine-readable parameter and provenance files. No figures were reproduced "
    "from external publications. Figure 1 is an author-created conceptual workflow "
    "and contains no independent quantitative data."
)

FIGURES = [
    (1, "figure01_model_overview", "Corrected Model Overview"),
    (2, "figure02_benchmark_trajectories", "Benchmark Encounter Trajectories"),
    (3, "figure03_initial_state_comparison", "Initial State Comparison"),
    (4, "figure04_mixing_escape_heatmaps", "Mixing versus Escape Heatmaps"),
    (5, "figure05_mixing_relaxation_heatmaps", "Mixing versus Relaxation Heatmaps"),
    (6, "figure06_reaction_selectivity", "Reaction Selectivity Sensitivity"),
    (7, "figure07_controls", "Controls and Limiting Cases"),
    (8, "figure08_solver_validation", "Independent Numerical Validation"),
    (9, "figure09_downstream_kinetics", "Downstream ROS Kinetics"),
    (10, "figure10_bulk_rate_comparison", "Literature Bulk Rate Comparison"),
    (11, "figure11_circuit_validation", "Three Qubit Coherent Embedding Validation"),
]

TABLES = [
    (1, "table01_original_versus_corrected_model", "Original versus Corrected ROS Spin Model"),
    (2, "table02_parameter_provenance", "Parameter Provenance"),
    (3, "table03_unavailable_parameters_and_consequences", "Unavailable Parameters and Consequences"),
    (4, "table04_baseline_results", "Baseline Results"),
    (5, "table05_controls_and_extrema", "Controls and Extrema"),
    (6, "table06_independent_solver_validation", "Independent Solver Validation"),
    (7, "table07_circuit_validation", "Circuit Validation"),
    (8, "table08_experimental_validation_evidence", "Experimental Validation Evidence"),
    (9, "table09_requirements_traceability", "Requirements Traceability"),
]

EQUATIONS = (
    "projectors",
    "hamiltonian",
    "master_equation",
    "reaction_yields",
    "probability_balance",
    "h2o2_stoichiometry",
)

SUPPORTED_CLAIMS = [
    "The spin-1/2 semiquinone and spin-1 ground-state oxygen electronic space has dimension six and decomposes into doublet and quartet manifolds of dimensions two and four.",
    "The implemented projectors, density matrices, Hamiltonian construction, trace-decreasing evolution, and probability accounting satisfy the repository's algebraic and numerical tests.",
    "Under declared dimensionless scenarios, primary-superoxide yield depends on competition among spin-selective electron transfer, doublet-quartet mixing, relaxation, and encounter escape.",
    "The adaptive Dormand-Prince implementation agrees with the constant-generator matrix-exponential reference within the declared numerical tolerances.",
    "The three-qubit calculation checks coherent six-state-to-eight-state encoding and simulator consistency when Qiskit executes.",
    "Condition-specific literature bulk rates and downstream aqueous kinetics can provide bounded context when their units, species, and experimental conditions remain explicit.",
]

PROHIBITED_CLAIMS = [
    "Absolute cellular superoxide or absolute in-vivo H2O2 concentration",
    "Cardiotoxicity or therapeutic response",
    "A measured kQ/kD value, ordering, probability distribution, or physical range",
    "A confirmed coherent or spin-correlated semiquinone-oxygen encounter",
    "A demonstrated doublet, quartet, entangled, or otherwise prepared encounter state",
    "A quantitative magnetic-field effect",
    "Transfer of bulk M^-1 s^-1 rates into encounter-level kD or kQ in s^-1",
    "Quantum advantage, hardware performance, or independent validation of the chemistry",
]

SCRIPT_MAPPING = [
    ("spin_chemistry.py", "Defines spin operators and projectors, six-state Hamiltonians, reference and independent propagators, staged ROS chemistry, evidence gates, and coherent embedding checks."),
    ("ROS.py", "Provides the guarded command line for a single evidence-backed bulk calculation or an explicitly opted-in encounter sensitivity case."),
    ("sensitivity_analysis.py", "Defines normalized one- and two-dimensional sweeps, benchmark initial states, probability outputs, and the non-predictive labeling policy."),
    ("paper_analysis.py", "Runs the deterministic paper workflow, writes Figure 1-11 source data, renders all figures, generates Tables 1-9, and records run metadata."),
    ("plotting.py", "Renders each figure from its saved source CSV using one visual style and the required sensitivity warning."),
    ("configs/doxorubicin_parameters.json", "Stores active, disabled, sensitivity-only, downstream, bulk-rate, and experimental-evidence records."),
    ("configs/parameter_provenance.csv", "Provides the machine-readable source ledger, conditions, uncertainty, suitability, and exact code paths."),
    ("tests", "Checks spin algebra, dynamics, evidence policy, independent validation, downstream balances, circuit behavior, legacy boundaries, and the paper and handoff pipelines."),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        raise ValueError(f"empty generated table: {path}")
    return rows


def equation_directory(results: Path) -> Path:
    candidates = (results / "metadata" / "equations", ROOT / "assets" / "equations")
    for candidate in candidates:
        if all((candidate / f"{name}.png").is_file() for name in EQUATIONS):
            return candidate
    raise FileNotFoundError(
        "typeset equation assets are missing; run "
        ".venv/bin/python render_equations.py assets/equations"
    )


def copy_equations_into_run(results: Path) -> None:
    target = results / "metadata" / "equations"
    if all((target / f"{name}.png").is_file() for name in EQUATIONS):
        return
    source = equation_directory(results)
    target.mkdir(parents=True, exist_ok=True)
    for name in EQUATIONS:
        shutil.copy2(source / f"{name}.png", target / f"{name}.png")


def required_files(results: Path) -> None:
    for _number, stem, _title in FIGURES:
        for suffix in ("png", "svg", "pdf"):
            path = results / "figures" / f"{stem}.{suffix}"
            if not path.is_file():
                raise FileNotFoundError(path)
        if not (results / "data" / f"{stem}.csv").is_file():
            raise FileNotFoundError(results / "data" / f"{stem}.csv")
        if not (results / "captions" / f"{stem}_caption.txt").is_file():
            raise FileNotFoundError(results / "captions" / f"{stem}_caption.txt")
    for _number, stem, _title in TABLES:
        for suffix in ("csv", "md"):
            path = results / "tables" / f"{stem}.{suffix}"
            if not path.is_file():
                raise FileNotFoundError(path)
    for relative in (
        "metadata/run_manifest.json",
        "metadata/paper_results_summary.json",
        "README.md",
    ):
        if not (results / relative).is_file():
            raise FileNotFoundError(results / relative)
    equation_directory(results)


def computed_findings(results: Path) -> dict:
    return json.loads(
        (results / "metadata" / "paper_results_summary.json").read_text(
            encoding="utf-8"
        )
    )


def summary_sections(results: Path) -> list[tuple[str, list[str]]]:
    finding = computed_findings(results)
    baseline_low, baseline_high = finding[
        "baseline_primary_superoxide_yield_range_per_encounter"
    ]
    grid_low, grid_high = finding[
        "selected_grid_primary_superoxide_yield_range_per_encounter"
    ]
    circuit_status = finding["circuit_status"]
    return [
        ("Original project objective", [
            "ROS_Spin originally asked whether electron-spin dynamics could influence reactive oxygen species formation during anthracycline redox cycling while retaining a quantum-circuit component for computational comparison. The completed workflow preserves that objective but replaces the original two-qubit chemical interpretation with the correct spin space for doxorubicin semiquinone and molecular oxygen.",
        ]),
        ("Scientific problems identified in the original implementation", [
            "The prior implementation used singlet/triplet and Bell-state ideas for a system whose oxygen partner has spin one. It also risked treating computational-basis parity and repeated circuit gates as chemical dynamics, carrying semiconductor relaxation data into an unrelated molecular encounter, and mapping spin outcomes directly to ROS products. Those choices could not support the proposed chemistry.",
        ]),
        ("Physical correction to spin one half times spin one", [
            "Doxorubicin or adriamycin semiquinone has electron spin S=1/2, whereas ground-state O2 has S=1. Their product space therefore contains 2 x 3 = 6 electronic states. This change determines the correct state dimension and invalidates a two-qubit singlet/triplet description of the encounter.",
        ]),
        ("Six state doublet quartet framework", [
            "Angular-momentum addition splits the six-state space into a doublet subspace of dimension two and a quartet subspace of dimension four. Exact projectors define the manifold observables. I6/6, PD/2, PQ/4, and bounded pD mixtures are computational benchmarks, not evidence of chemical state preparation.",
        ]),
        ("Hamiltonian and density matrix evolution", [
            "The active model propagates a six-by-six density matrix with a constant-generator matrix exponential. Zeeman, exchange, dipolar, oxygen zero-field-splitting, and local-field structures are implemented, but unavailable exact-system terms remain disabled or are activated only as named sensitivity coordinates. Local isotropic Lindblad terms supply an explicitly phenomenological relaxation sensitivity model.",
        ]),
        ("Spin selective electron transfer and escape", [
            "Electron transfer is represented by separate first-order loss operators in the doublet and quartet manifolds, while encounter escape removes population independently. The model integrates these fluxes through time and retains unresolved survival so that reaction, escape, and survival close to unity.",
        ]),
        ("Primary superoxide calculation", [
            "Primary superoxide is calculated only as the sum of integrated doublet and quartet electron-transfer yields. One reacted encounter contributes one primary superoxide-family radical equivalent. A final spin population is never re-labeled as chemical product.",
        ]),
        ("Separate downstream hydrogen peroxide kinetics", [
            "A distinct fixed-pH model partitions the radical pool between HO2 and O2-minus, applies spontaneous and SOD-mediated dismutation, consumes two radical equivalents per H2O2, and tracks an optional H2O2 loss channel. This stage does not feed an encounter spin population directly into H2O2.",
        ]),
        ("Evidence gated parameter policy", [
            "The JSON configuration and provenance CSV distinguish measured or literature-derived bulk quantities from missing encounter quantities. Bulk constants in M^-1 s^-1 remain separate from kD, kQ, relaxation, and escape in s^-1. Missing geometry, tensors, relaxation times, state preparation, and state-resolved rates are not assigned invented values.",
        ]),
        ("Sensitivity analysis design", [
            "All encounter results are conditional, dimensionless sensitivity results. Mixing, relaxation, escape, initial doublet fraction, and kQ/kD are varied over declared computational coordinates. The equality kQ/kD=1 is the spin-independent null; the displayed axes are not measured ranges, probability distributions, priors, or confidence intervals.",
        ]),
        ("Main computed findings", [
            f"Across the five baseline initial-state benchmarks, the computed primary-superoxide yield per encounter spans {baseline_low:.6g} to {baseline_high:.6g}. Across the selected sensitivity grids it spans {grid_low:.6g} to {grid_high:.6g}; these extrema describe only the executed grid. The simulations demonstrate how primary superoxide yield would depend on the competition among spin-selective electron transfer, doublet-quartet mixing, relaxation, and encounter escape under specified dimensionless scenarios. They do not establish that the assumed spin preparation, mixing magnitude, or state-selective rates occur in the doxorubicin semiquinone-oxygen system.",
            "Computed structural results include projector algebra, six-state dimensions, valid density matrices, probability accounting, solver agreement, and coherent-embedding consistency. Conditional sensitivity results cover mixing, relaxation, escape, initial state, and kQ/kD. Literature-derived calculations remain confined to condition-specific bulk rates, downstream aqueous kinetics, and provenance-qualified experimental context.",
        ]),
        ("Numerical validation", [
            f"The worst absolute discrepancy between the reference matrix exponential and the separately constructed adaptive Dormand-Prince solver is {finding['worst_independent_solver_absolute_error']:.3e}. The maximum encounter probability-balance error is {finding['maximum_encounter_probability_balance_error']:.3e}, and the maximum downstream radical-equivalent balance error is {finding['maximum_downstream_radical_equivalent_balance_error_m']:.3e} M. These checks validate numerical consistency rather than physical parameterization.",
        ]),
        ("Quantum circuit validation", [
            f"Circuit status for the recorded run is {circuit_status}. The three-qubit calculation embeds the six physical states within an eight-state space, checks the two unused states for leakage, and compares statevectors and doublet/quartet observables. It validates coherent encoding consistency only; it does not demonstrate quantum advantage or independently validate the chemistry.",
        ]),
        ("Relationship to anthracycline ROS biology", [
            "The corrected workflow keeps the biological motivation at the level supported by current evidence: anthracycline semiquinone chemistry can transfer an electron to oxygen, and downstream superoxide chemistry can form H2O2. The present encounter model does not bridge its illustrative per-encounter yields to organelle, cell, animal, or patient exposure without missing formation, association, residence, competing-sink, transport, and compartment parameters.",
        ]),
        ("Supported conclusions", [
            "The project now supports structural statements about the correct spin space, conditional dependence within the declared equations, numerical solver agreement, coherent embedding consistency, and condition-specific literature arithmetic. The complete supported-claims list appears in the handoff report.",
        ]),
        ("Unsupported conclusions", [
            "The project does not quantitatively predict absolute cellular superoxide, in-vivo H2O2, cardiotoxicity, therapeutic response, a measured kQ/kD, a confirmed spin-correlated encounter, a magnetic-field effect, or quantum advantage. The complete prohibited-claims list appears in the handoff report.",
        ]),
        ("Recommended framing for the rewritten paper", [
            "Frame the work as an evidence-gated computational framework and conditional sensitivity study connected to the original anthracycline ROS question. Lead with the physical correction, present computed structural validation separately from sensitivity results and literature-derived calculations, and state the missing measurements as requirements for future quantitative chemistry rather than filling them with analogues.",
        ]),
    ]


def write_summary(results: Path, output: Path) -> None:
    lines = [
        "# ROS Spin Professor Summary",
        "",
        PROVENANCE_STATEMENT,
        "",
    ]
    for index, (heading, paragraphs) in enumerate(summary_sections(results), 1):
        lines.extend([f"## {index} {heading}", ""])
        for paragraph in paragraphs:
            lines.extend([paragraph, ""])
    lines.extend(["## Active script contributions", ""])
    for name, meaning in SCRIPT_MAPPING:
        lines.append(f"- `{name}`: {meaning}")
    lines.extend(["", "## Supported claims", ""])
    lines.extend(f"- {claim}" for claim in SUPPORTED_CLAIMS)
    lines.extend(["", "## Prohibited claims", ""])
    lines.extend(f"- {claim}" for claim in PROHIBITED_CLAIMS)
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.first_child_found_in("w:shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def compact_value(value: str) -> str:
    return str(value or "")


def add_generated_table(
    doc: Document,
    rows: list[dict[str, str]],
    fields: list[tuple[str, str]],
    *,
    font_size: float = 7.5,
) -> None:
    table = doc.add_table(rows=1, cols=len(fields))
    table.autofit = True
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (_field, label) in enumerate(fields):
        cell = header.cells[index]
        cell.text = label
        shade_cell(cell, "1F4E78")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
    for row_index, source in enumerate(rows):
        cells = table.add_row().cells
        for index, (field, _label) in enumerate(fields):
            cell = cells[index]
            cell.text = compact_value(source.get(field, ""))
            if row_index % 2:
                shade_cell(cell, "EAF2F8")
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    doc.add_paragraph()


def add_table_source(doc: Document, results: Path, stem: str) -> None:
    paragraph = doc.add_paragraph(style="Source Note")
    paragraph.add_run("Generated CSV source  ").bold = True
    paragraph.add_run(str((results / "tables" / f"{stem}.csv").resolve()))


def add_equation_image(doc: Document, path: Path) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(path) as image:
        width_px, height_px = image.size
    width = min(6.25, max(2.5, width_px / 500))
    if width / (width_px / height_px) > 0.72:
        width = 0.72 * (width_px / height_px)
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_figure(doc: Document, results: Path, number: int, stem: str, title: str) -> None:
    image_path = results / "figures" / f"{stem}.png"
    caption_path = results / "captions" / f"{stem}_caption.txt"
    caption = caption_path.read_text(encoding="utf-8").strip()
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    max_width, max_height = 6.65, 5.55
    width = min(max_width, max_height * aspect)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(style="Figure Caption")
    cap.paragraph_format.keep_together = True
    run = cap.add_run(f"Figure {number}  {title}. ")
    run.bold = True
    caption_body = caption.split(". ", 1)[1] if ". " in caption else caption
    cap.add_run(caption_body)
    source = doc.add_paragraph(style="Source Note")
    source.paragraph_format.keep_together = True
    source.add_run("Source data  ").bold = True
    source.add_run(str((results / "data" / f"{stem}.csv").resolve()))
    source.add_run("    Caption  ").bold = True
    source.add_run(str(caption_path.resolve()))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    for node in (begin, instruction, separate, text, end):
        run.append(node)
    paragraph._p.append(run)


def set_section_geometry(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)


def apply_document_style(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_properties = title.element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)
    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    if "Figure Caption" not in [style.name for style in styles]:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(3)
    if "Source Note" not in [style.name for style in styles]:
        source = styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source = styles["Source Note"]
    source.font.name = "Aptos"
    source.font.size = Pt(7.5)
    source.font.color.rgb = RGBColor(85, 85, 85)
    source.paragraph_format.space_after = Pt(6)


def add_header_footer(section, first_page: bool = False) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.different_first_page_header_footer = first_page
    if not first_page:
        header = section.header.paragraphs[0]
        header.clear()
        header.text = "ROS Spin Professor Handoff"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 90, 90)
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)


def add_section_heading(doc: Document, title: str, *, page_break: bool = False) -> None:
    if page_break:
        doc.add_page_break()
    doc.add_heading(title, level=1)


def add_paragraphs(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        doc.add_paragraph(text)


def section_lookup(results: Path) -> dict[str, list[str]]:
    return {heading: paragraphs for heading, paragraphs in summary_sections(results)}


def table_rows(results: Path, stem: str) -> list[dict[str, str]]:
    return read_csv(results / "tables" / f"{stem}.csv")


def add_table_one(doc: Document, results: Path) -> None:
    stem = TABLES[0][1]
    add_generated_table(
        doc,
        table_rows(results, stem),
        [
            ("model_element", "Model element"),
            ("original_implementation", "Original implementation"),
            ("corrected_implementation", "Corrected implementation"),
            ("scientific_consequence", "Consequence"),
        ],
        font_size=8,
    )
    add_table_source(doc, results, stem)


def add_results_tables(doc: Document, results: Path) -> None:
    definitions = [
        (4, [
            ("initial_state_definition", "Initial state"),
            ("p_doublet_initial", "Initial pD"),
            ("doublet_reaction_yield_per_encounter", "D reaction"),
            ("quartet_reaction_yield_per_encounter", "Q reaction"),
            ("primary_superoxide_yield_per_encounter", "Primary superoxide"),
            ("escape_yield_per_encounter", "Escape"),
            ("survival_probability", "Survival"),
            ("probability_balance", "Balance"),
        ]),
        (5, [
            ("summary_type", "Type"), ("source_grid", "Source"),
            ("control_label", "Control"), ("initial_state_definition", "Initial state"),
            ("mixing_over_reference", "Mixing/kref"),
            ("radical_relaxation_over_reference", "Relaxation/kref"),
            ("escape_over_reference", "Escape/kref"), ("kq_over_kd", "kQ/kD"),
            ("primary_superoxide_yield_per_encounter", "Primary superoxide"),
            ("escape_yield_per_encounter", "Escape yield"),
            ("unresolved_probability", "Survival"),
        ]),
        (6, [
            ("observable", "Observable"), ("worst_absolute_error", "Worst absolute error"),
            ("worst_relative_error", "Worst relative error"),
            ("declared_absolute_tolerance", "Tolerance"), ("passes", "Pass"),
            ("reference_solver", "Reference solver"),
            ("independent_solver", "Independent solver"), ("scope", "Scope"),
        ]),
    ]
    for number, fields in definitions:
        stem, title = TABLES[number - 1][1], TABLES[number - 1][2]
        doc.add_heading(f"Table {number}  {title}", level=2)
        add_generated_table(doc, table_rows(results, stem), fields, font_size=7.1)
        add_table_source(doc, results, stem)
    number = 7
    stem, title = TABLES[number - 1][1], TABLES[number - 1][2]
    rows = table_rows(results, stem)
    doc.add_heading(f"Table {number}  {title}", level=2)
    doc.add_heading("Table 7A Numerical metrics", level=3)
    add_generated_table(doc, rows, [
        ("metric", "Metric"), ("validation_layer", "Layer"),
        ("value", "Value"), ("tolerance", "Tolerance"), ("passes", "Pass"),
    ], font_size=7.5)
    doc.add_heading("Table 7B Execution status and scope", level=3)
    add_generated_table(doc, rows, [
        ("metric", "Metric"), ("qiskit_available", "Qiskit available"),
        ("qiskit_execution_status", "Execution"), ("implementation_type", "Implementation"),
        ("scope", "Scope"), ("reason", "Reason"),
    ], font_size=7.3)
    add_table_source(doc, results, stem)


def add_parameter_table(doc: Document, results: Path) -> None:
    number = 2
    stem, title = TABLES[number - 1][1], TABLES[number - 1][2]
    rows = table_rows(results, stem)
    doc.add_heading(f"Table {number}  {title}", level=1)
    doc.add_paragraph("The table is split into two aligned panels for print legibility. The symbol column links the panels; the generated CSV preserves the complete row schema.")
    doc.add_heading("Table 2A Identity conditions and value", level=2)
    add_generated_table(doc, rows, [
        ("symbol", "Symbol"), ("definition", "Definition"),
        ("value_or_range", "Value or range"), ("unit", "Unit"),
        ("species", "Exact species"), ("charge_state", "Charge or protonation"),
        ("environment", "Environment"), ("temperature", "Temperature"), ("pH", "pH"),
    ], font_size=6.4)
    doc.add_heading("Table 2B Method provenance and suitability", level=2)
    add_generated_table(doc, rows, [
        ("symbol", "Symbol"), ("method", "Method"), ("source", "DOI or direct source"),
        ("status", "Status"), ("uncertainty", "Uncertainty"),
        ("limitations", "Limitations"), ("permitted_use", "Suitability"),
        ("code_location", "Code location"),
    ], font_size=6.4)
    add_table_source(doc, results, stem)


def add_unavailable_table(doc: Document, results: Path) -> None:
    number = 3
    stem, title = TABLES[number - 1][1], TABLES[number - 1][2]
    doc.add_heading(f"Table {number}  {title}", level=1)
    add_generated_table(doc, table_rows(results, stem), [
        ("missing_parameter", "Unavailable parameter"), ("why_needed", "Why needed"),
        ("sensitivity_coordinate_allowed", "Sensitivity allowed"),
        ("range_status", "Range status"), ("allowed_treatment", "Allowed treatment"),
        ("conclusion_prohibited", "Prohibited conclusion"),
    ], font_size=7.1)
    add_table_source(doc, results, stem)


def add_experimental_table(doc: Document, results: Path) -> None:
    number = 8
    stem, title = TABLES[number - 1][1], TABLES[number - 1][2]
    rows = table_rows(results, stem)
    doc.add_heading(f"Table {number}  {title}", level=1)
    doc.add_heading("Table 8A System conditions and reported result", level=2)
    add_generated_table(doc, rows, [
        ("dataset_id", "Record"), ("exact_chemical_or_biological_system", "Exact system"),
        ("observable", "Observable"), ("conditions", "Conditions"),
        ("reported_result", "Reported result"), ("unit", "Unit"),
    ], font_size=6.8)
    doc.add_heading("Table 8B Evidence method and limitation", level=2)
    add_generated_table(doc, rows, [
        ("dataset_id", "Record"), ("method", "Method"), ("source", "Source"),
        ("usable_for_direct_validation", "Direct validation"),
        ("limitations", "Limitation"),
    ], font_size=7.0)
    add_table_source(doc, results, stem)


def add_traceability_table(doc: Document, results: Path) -> None:
    number = 9
    stem, title = TABLES[number - 1][1], TABLES[number - 1][2]
    rows = table_rows(results, stem)
    doc.add_heading(f"Table {number}  {title}", level=1)
    doc.add_heading("Table 9A Requirement to implementation mapping", level=2)
    add_generated_table(doc, rows, [
        ("requirement_id", "Requirement"), ("research_topic_or_output", "Topic or output"),
        ("exact_file", "Exact file"), ("exact_code_object", "Exact code object"),
        ("figure_or_table", "Figure or table"),
    ], font_size=7.0)
    doc.add_heading("Table 9B Test evidence and remaining limitation", level=2)
    add_generated_table(doc, rows, [
        ("requirement_id", "Requirement"), ("test", "Test"),
        ("evidence_status", "Evidence status"),
        ("remaining_limitation", "Remaining limitation"),
    ], font_size=7.2)
    add_table_source(doc, results, stem)


def extract_bibliography() -> list[str]:
    source = ROOT / "references" / "doxorubicin_parameter_review.md"
    lines = source.read_text(encoding="utf-8").splitlines()
    try:
        start = lines.index("## Complete bibliography") + 1
    except ValueError:
        return []
    bibliography = []
    for line in lines[start:]:
        if line.startswith("## "):
            break
        if line.startswith("- "):
            bibliography.append(line[2:].replace("*", ""))
    return bibliography


def build_docx(results: Path, output: Path) -> None:
    sections = section_lookup(results)
    manifest = json.loads((results / "metadata" / "run_manifest.json").read_text())
    doc = Document()
    apply_document_style(doc)
    set_section_geometry(doc.sections[0], landscape=False)
    add_header_footer(doc.sections[0], first_page=True)
    properties = doc.core_properties
    properties.title = TITLE
    properties.subject = SUBTITLE
    properties.author = "ROS_Spin project team"
    properties.keywords = "doxorubicin semiquinone oxygen doublet quartet ROS"

    top = doc.add_paragraph(style="Title")
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    top.add_run(TITLE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(SUBTITLE)
    run.bold = True
    run.font.size = Pt(15)
    doc.add_paragraph()
    identity = doc.add_paragraph()
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    identity.add_run(f"Repository commit {manifest['environment']['git_commit'][:12]}\n")
    identity.add_run(f"Generated {manifest['timestamp_iso8601']}\n")
    identity.add_run("Professor-facing final computational report")
    doc.add_paragraph()
    provenance = doc.add_paragraph(PROVENANCE_STATEMENT)
    provenance.alignment = WD_ALIGN_PARAGRAPH.CENTER
    provenance.paragraph_format.space_before = Pt(24)
    provenance.paragraph_format.space_after = Pt(24)
    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope_run = scope.add_run("NON-PREDICTIVE DIMENSIONLESS SENSITIVITY ANALYSIS")
    scope_run.bold = True
    scope_run.font.color.rgb = RGBColor(139, 26, 26)

    add_section_heading(doc, "Purpose of This Handoff", page_break=True)
    doc.add_paragraph(
        "This report hands off the completed corrected ROS_Spin implementation, its generated scientific results, numerical validation, evidence record, and reproduction package. It is written to support technical review and manuscript revision without requiring the reader to reconstruct the repository history."
    )
    doc.add_paragraph(PROVENANCE_STATEMENT)

    add_section_heading(doc, "Executive Summary")
    add_paragraphs(doc, sections["Main computed findings"])
    add_paragraphs(doc, sections["Numerical validation"])
    add_paragraphs(doc, sections["Quantum circuit validation"])

    add_section_heading(doc, "Connection to the Original ROS Spin Project")
    add_paragraphs(doc, sections["Original project objective"])
    add_paragraphs(doc, sections["Scientific problems identified in the original implementation"])
    add_paragraphs(doc, sections["Relationship to anthracycline ROS biology"])

    add_section_heading(doc, "Table 1 Original versus Corrected Implementation", page_break=True)
    doc.add_paragraph("Table 1 connects the corrected implementation to the original project and identifies why each change matters scientifically.")
    add_table_one(doc, results)

    add_section_heading(doc, "Corrected Model and Equations", page_break=True)
    add_paragraphs(doc, sections["Physical correction to spin one half times spin one"])
    add_paragraphs(doc, sections["Six state doublet quartet framework"])
    equation_dir = equation_directory(results)
    add_equation_image(doc, equation_dir / "projectors.png")
    add_paragraphs(doc, sections["Hamiltonian and density matrix evolution"])
    add_equation_image(doc, equation_dir / "hamiltonian.png")
    add_paragraphs(doc, sections["Spin selective electron transfer and escape"])
    add_equation_image(doc, equation_dir / "master_equation.png")
    add_paragraphs(doc, sections["Primary superoxide calculation"])
    add_equation_image(doc, equation_dir / "reaction_yields.png")
    add_equation_image(doc, equation_dir / "probability_balance.png")
    add_paragraphs(doc, sections["Separate downstream hydrogen peroxide kinetics"])
    add_equation_image(doc, equation_dir / "h2o2_stoichiometry.png")

    add_section_heading(doc, "Figure 1 Conceptual Model Overview", page_break=True)
    add_figure(doc, results, *FIGURES[0])

    add_section_heading(doc, "Computational Methods and Script Mapping", page_break=True)
    doc.add_paragraph("The active workflow separates model equations, command-line policy, sensitivity design, plotting, evidence, and tests so that each generated claim can be traced to code and source data.")
    for name, meaning in SCRIPT_MAPPING:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{name}  ").bold = True
        paragraph.add_run(meaning)
    add_paragraphs(doc, sections["Evidence gated parameter policy"])
    add_paragraphs(doc, sections["Sensitivity analysis design"])

    figure_section_titles = {
        2: "Figure 2 Benchmark Trajectories",
        3: "Figure 3 Initial State Comparison",
        4: "Figures 4 and 5 Timescale Competition",
        6: "Figure 6 Reaction Selectivity Sensitivity",
        7: "Figure 7 Controls",
        8: "Figure 8 Independent Solver Validation",
        9: "Figure 9 Downstream ROS Kinetics",
        10: "Figure 10 Primary Source Bulk Rate Context",
        11: "Figure 11 Circuit Embedding Validation",
    }
    for number, stem, title in FIGURES[1:]:
        if number == 5:
            doc.add_page_break()
        else:
            add_section_heading(doc, figure_section_titles[number], page_break=True)
        if number == 8:
            add_paragraphs(doc, sections["Numerical validation"])
        elif number == 9:
            add_paragraphs(doc, sections["Separate downstream hydrogen peroxide kinetics"])
        elif number == 10:
            doc.add_paragraph("The plotted literature values remain condition-specific bulk kinetic records. Their units and incomplete conditions prevent their use as encounter-level state-resolved reaction rates.")
        elif number == 11:
            add_paragraphs(doc, sections["Quantum circuit validation"])
        add_figure(doc, results, number, stem, title)

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(landscape, landscape=True)
    add_header_footer(landscape)
    doc.add_heading("Baseline and Control Tables", level=1)
    doc.add_paragraph("Tables 4-7 report the generated baseline, controls, independent-solver errors, and coherent-embedding metrics. All numeric cells are populated from their CSV sources.")
    add_results_tables(doc, results)
    add_parameter_table(doc, results)
    add_unavailable_table(doc, results)
    add_experimental_table(doc, results)
    add_traceability_table(doc, results)

    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(portrait, landscape=False)
    add_header_footer(portrait)
    doc.add_heading("Supported Claims", level=1)
    add_bullets(doc, SUPPORTED_CLAIMS)
    doc.add_heading("Unsupported and Prohibited Claims", level=1)
    add_bullets(doc, PROHIBITED_CLAIMS)
    doc.add_heading("Recommended Manuscript Framing", level=1)
    add_paragraphs(doc, sections["Recommended framing for the rewritten paper"])
    doc.add_paragraph(
        "Recommended manuscript language: The simulations demonstrate how primary superoxide yield would depend on the competition among spin-selective electron transfer, doublet-quartet mixing, relaxation, and encounter escape under specified dimensionless scenarios. They do not establish that the assumed spin preparation, mixing magnitude, or state-selective rates occur in the doxorubicin semiquinone-oxygen system."
    )
    doc.add_heading("Exact Reproduction Instructions", level=1)
    doc.add_paragraph("From the repository root, run the following commands after installing the pinned requirements and ensuring LibreOffice is available:")
    command = doc.add_paragraph()
    command.style = doc.styles["Normal"]
    run = command.add_run(
        ".venv/bin/python paper_analysis.py --mode full --output-dir results/paper --execute-circuit --overwrite\n"
        ".venv/bin/python render_equations.py assets/equations\n"
        "python3 professor_handoff.py --results-dir results/paper --deliverables-dir deliverables --overwrite"
    )
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    doc.add_paragraph("The exact numerical invocation, dependency versions, source hashes, assumptions, tolerances, Qiskit status, and limitations are recorded in results/paper/metadata/run_manifest.json.")

    doc.add_heading("References and DOI List", level=1)
    bibliography = extract_bibliography()
    if not bibliography:
        doc.add_paragraph("See references/doxorubicin_parameter_review.md.")
    else:
        for entry in bibliography:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(entry)

    doc.add_heading("Appendix Figure Captions and Table Notes", level=1)
    doc.add_paragraph("Each figure caption below is copied from the generated caption file that accompanies the plotted source data.")
    for number, stem, title in FIGURES:
        doc.add_heading(f"Figure {number} {title}", level=2)
        doc.add_paragraph((results / "captions" / f"{stem}_caption.txt").read_text(encoding="utf-8").strip())
        add_table_source_note = doc.add_paragraph(style="Source Note")
        add_table_source_note.add_run("Figure source  ").bold = True
        add_table_source_note.add_run(str((results / "data" / f"{stem}.csv").resolve()))
    for number, stem, title in TABLES:
        doc.add_heading(f"Table {number} {title}", level=2)
        doc.add_paragraph("Generated programmatically from repository code or machine-readable evidence. Interpret all values within the status, conditions, suitability, and limitations columns.")
        add_table_source(doc, results, stem)

    doc.save(output)


def convert_to_pdf(docx_path: Path, output_dir: Path, soffice: str | None) -> Path:
    executable = soffice or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("LibreOffice/soffice is unavailable; cannot create the matched PDF")
    completed = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
        text=True,
        capture_output=True,
    )
    if completed.returncode:
        raise RuntimeError(f"LibreOffice conversion failed: {completed.stderr or completed.stdout}")
    pdf = output_dir / f"{docx_path.stem}.pdf"
    if not pdf.is_file() or pdf.stat().st_size < 1000:
        raise RuntimeError(f"LibreOffice did not create a valid PDF: {completed.stdout}")
    return pdf


def update_result_metadata(
    results: Path,
    deliverables: Path,
    *,
    visual_qa_status: str,
    command: str,
) -> None:
    manifest_path = results / "metadata" / "run_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    named = {
        "professor_summary": deliverables / "ROS_Spin_professor_summary.md",
        "word_document": deliverables / "ROS_Spin_professor_handoff.docx",
        "pdf_document": deliverables / "ROS_Spin_professor_handoff.pdf",
    }
    manifest["professor_handoff"] = {
        "generation_command": command,
        "updated_utc": datetime.now(timezone.utc).isoformat(),
        "visual_qa_status": visual_qa_status,
        "artifacts": {
            key: {
                "path": str(path.resolve()),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
            for key, path in named.items()
        },
        "zip_checksum_note": "The ZIP hash is written to a sidecar after archive creation to avoid self-reference.",
        "source_hashes_sha256": {
            "professor_handoff.py": sha256(ROOT / "professor_handoff.py"),
            "render_equations.py": sha256(ROOT / "render_equations.py"),
            "requirements.txt": sha256(ROOT / "requirements.txt"),
        },
    }
    excluded_hash_targets = {
        manifest_path,
        results / "metadata" / "output_checksums.csv",
        results / "metadata" / "SHA256SUMS",
    }
    hashed_files = sorted(
        path for path in results.rglob("*")
        if path.is_file() and path not in excluded_hash_targets
    )
    manifest["generated_file_hashes_sha256"] = {
        str(path.relative_to(results)): sha256(path) for path in hashed_files
    }
    manifest["generated_files"] = sorted(
        {str(path.relative_to(results)) for path in results.rglob("*") if path.is_file()}
        | {"metadata/run_manifest.json", "metadata/output_checksums.csv", "metadata/SHA256SUMS"}
    )
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    checksum_path = results / "metadata" / "output_checksums.csv"
    checksum_rows = []
    for path in sorted(path for path in results.rglob("*") if path.is_file()):
        if path in (checksum_path, results / "metadata" / "SHA256SUMS"):
            continue
        relative = str(path.relative_to(results))
        checksum_rows.append({
            "path": relative,
            "sha256": sha256(path),
            "bytes": path.stat().st_size,
            "artifact_kind": relative.split("/", 1)[0] if "/" in relative else "run_root",
        })
    with checksum_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(checksum_rows[0]))
        writer.writeheader()
        writer.writerows(checksum_rows)
    sha_path = results / "metadata" / "SHA256SUMS"
    lines = [
        f"{sha256(path)}  {path.relative_to(results)}"
        for path in sorted(path for path in results.rglob("*") if path.is_file())
        if path != sha_path
    ]
    sha_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def update_reproduction_readme(results: Path) -> None:
    readme = results / "README.md"
    text = readme.read_text(encoding="utf-8")
    marker = "## Professor handoff reproduction"
    if marker in text:
        text = text.split(marker, 1)[0].rstrip() + "\n\n"
    text += (
        f"{marker}\n\n"
        "After generating the full paper results, create the professor package with:\n\n"
        "```bash\n"
        ".venv/bin/python render_equations.py assets/equations\n"
        "python3 professor_handoff.py --results-dir results/paper --deliverables-dir deliverables --overwrite\n"
        "```\n\n"
        "The document step requires python-docx and LibreOffice. The ZIP contains the DOCX, matched PDF, summary, all figure formats and source data, Tables 1-9, captions, metadata, and this README.\n"
    )
    readme.write_text(text, encoding="utf-8")


def build_zip(results: Path, deliverables: Path) -> Path:
    zip_path = deliverables / "ROS_Spin_professor_handoff.zip"
    root_name = "ROS_Spin_professor_handoff"
    members = [
        deliverables / "ROS_Spin_professor_handoff.docx",
        deliverables / "ROS_Spin_professor_handoff.pdf",
        deliverables / "ROS_Spin_professor_summary.md",
        results / "README.md",
    ]
    for folder in ("figures", "data", "tables", "captions", "metadata"):
        members.extend(sorted(path for path in (results / folder).rglob("*") if path.is_file()))
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in members:
            if path.parent == deliverables:
                relative = path.name
            elif path == results / "README.md":
                relative = "reproduction_README.md"
            else:
                relative = str(path.relative_to(results))
            archive.write(path, f"{root_name}/{relative}")
    sidecar = zip_path.with_suffix(".zip.sha256")
    sidecar.write_text(f"{sha256(zip_path)}  {zip_path.name}\n", encoding="utf-8")
    return zip_path


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-dir", type=Path, default=DEFAULT_RESULTS)
    parser.add_argument("--deliverables-dir", type=Path, default=DEFAULT_DELIVERABLES)
    parser.add_argument("--soffice")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--package-only", action="store_true")
    parser.add_argument("--visual-qa-status", default="pending", choices=("pending", "passed", "failed"))
    return parser


def main(argv: list[str] | None = None) -> None:
    args = build_parser().parse_args(argv)
    results = args.results_dir.resolve()
    deliverables = args.deliverables_dir.resolve()
    copy_equations_into_run(results)
    required_files(results)
    deliverables.mkdir(parents=True, exist_ok=True)
    docx = deliverables / "ROS_Spin_professor_handoff.docx"
    pdf = deliverables / "ROS_Spin_professor_handoff.pdf"
    summary = deliverables / "ROS_Spin_professor_summary.md"
    archive = deliverables / "ROS_Spin_professor_handoff.zip"
    targets = (docx, pdf, summary, archive, archive.with_suffix(".zip.sha256"))
    if not args.package_only:
        existing = [path for path in targets if path.exists()]
        if existing and not args.overwrite:
            raise FileExistsError(
                "deliverables already exist; use --overwrite: "
                + ", ".join(str(path) for path in existing)
            )
        if args.overwrite:
            for path in existing:
                path.unlink()
        write_summary(results, summary)
        build_docx(results, docx)
        convert_to_pdf(docx, deliverables, args.soffice)
    else:
        for path in (docx, pdf, summary):
            if not path.is_file():
                raise FileNotFoundError(path)
        for path in (archive, archive.with_suffix(".zip.sha256")):
            if path.exists():
                path.unlink()

    invocation = "python3 professor_handoff.py " + " ".join(
        f"{name} {value}" for name, value in (
            ("--results-dir", args.results_dir),
            ("--deliverables-dir", args.deliverables_dir),
        )
    )
    update_reproduction_readme(results)
    update_result_metadata(
        results,
        deliverables,
        visual_qa_status=args.visual_qa_status,
        command=invocation,
    )
    zip_path = build_zip(results, deliverables)
    print(json.dumps({
        "status": "complete",
        "docx": str(docx),
        "pdf": str(pdf),
        "summary": str(summary),
        "zip": str(zip_path),
        "zip_sha256": sha256(zip_path),
        "visual_qa_status": args.visual_qa_status,
    }, sort_keys=True))


if __name__ == "__main__":
    main()
