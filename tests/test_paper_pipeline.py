"""End-to-end and contract tests for the active paper-results pipeline."""
from __future__ import annotations

import csv
import hashlib
import importlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import numpy as np
from PIL import Image, ImageStat
import matplotlib.pyplot as plt
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sensitivity_analysis import encounter_result_row  # noqa: E402
from paper_analysis import TOLERANCES  # noqa: E402
from professor_handoff import PROVENANCE_STATEMENT, build_docx  # noqa: E402
from plotting import (  # noqa: E402
    plot_bulk_rates,
    plot_controls,
    plot_solver_validation,
)
from spin_chemistry import (  # noqa: E402
    ModelPolicyError,
    USE_CONDITION_MATCHED,
    evaluate_bulk_superoxide_rate,
    validate_authority_bundle,
)


class PaperPipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory()
        cls.output = Path(cls.temporary.name) / "paper"
        completed = subprocess.run(
            [
                sys.executable, "paper_analysis.py", "--mode", "quick",
                "--samples", "21", "--grid-size", "4",
                "--formats", "png", "--dpi", "100", "--seed", "1729",
                "--execute-circuit",
                "--output-dir", str(cls.output),
            ],
            cwd=ROOT, text=True, capture_output=True,
        )
        if completed.returncode:
            raise AssertionError(completed.stderr)
        cls.report = json.loads(completed.stdout)

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def _rows(self, stem: str, folder: str = "data"):
        with (self.output / folder / f"{stem}.csv").open(newline="") as handle:
            return list(csv.DictReader(handle))

    def test_active_modules_and_clis_import(self):
        for module in (
            "spin_chemistry", "ROS", "sensitivity_analysis", "plotting",
            "paper_analysis", "professor_handoff",
        ):
            imported = importlib.import_module(module)
            self.assertIsNotNone(imported, module)
        for script in (
            "ROS.py", "sensitivity_analysis.py", "paper_analysis.py",
            "professor_handoff.py",
        ):
            completed = subprocess.run(
                [sys.executable, script, "--help"], cwd=ROOT,
                text=True, capture_output=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)

    def test_active_files_do_not_reference_removed_two_qubit_api(self):
        active = (
            "spin_chemistry.py", "ROS.py", "sensitivity_analysis.py",
            "plotting.py", "paper_analysis.py", "professor_handoff.py",
        )
        removed = ("build_rp_circuit", "counts_to_ros", "coherent_circuit_validation", "f_ow", "f_fb")
        for filename in active:
            source = (ROOT / filename).read_text(encoding="utf-8")
            for symbol in removed:
                self.assertNotIn(symbol, source, f"{filename}: {symbol}")

    def test_legacy_entrypoints_are_unambiguous_and_sources_are_retained(self):
        scripts = (
            "ROS_Util.py", "basecheck.py", "error.py", "fig09plot.py",
            "heat_gamma.py", "heat_singlet.py", "looped.py", "noise.py", "plot_error.py",
        )
        for script in scripts:
            self.assertTrue((ROOT / "legacy" / script).is_file())
            completed = subprocess.run(
                [sys.executable, script], cwd=ROOT, text=True, capture_output=True,
            )
            self.assertNotEqual(completed.returncode, 0)
            self.assertIn("LEGACY SEMICONDUCTOR/TWO-QUBIT MATERIAL", completed.stderr)

    def test_quick_generation_has_complete_directory_and_artifact_contract(self):
        self.assertEqual(self.report["status"], "complete")
        for directory in ("data", "figures", "tables", "captions", "metadata"):
            self.assertTrue((self.output / directory).is_dir())
        self.assertTrue((self.output / "README.md").is_file())
        for number in range(1, 12):
            stem = f"figure{number:02d}_"
            sources = list((self.output / "data").glob(f"{stem}*.csv"))
            self.assertEqual(len(sources), 1, stem)
            self.assertGreater(sources[0].stat().st_size, 100)
            captions = list((self.output / "captions").glob(f"{stem}*_caption.txt"))
            self.assertEqual(len(captions), 1, stem)
        for number in range(1, 10):
            self.assertEqual(len(list((self.output / "tables").glob(f"table{number:02d}_*.csv"))), 1)
            self.assertEqual(len(list((self.output / "tables").glob(f"table{number:02d}_*.md"))), 1)

    def test_unavailable_parameter_consequences_are_parameter_specific(self):
        rows = {
            row["missing_parameter"]: row
            for row in self._rows(
                "table03_unavailable_parameters_and_consequences", "tables"
            )
        }
        self.assertIn("observation endpoint or residence time", rows["encounter duration"]["why_needed"])
        self.assertIn("solution encounter", rows["encounter-specific O2 relaxation"]["why_needed"])
        self.assertIn("removal of downstream H2O2", rows["compartment-specific H2O2 loss rate"]["why_needed"])
        self.assertIn("compartment-level H2O2", rows["compartment-specific H2O2 loss rate"]["conclusion_prohibited"])
        self.assertTrue((self.output / "metadata" / "output_checksums.csv").is_file())

    def test_every_rendered_figure_is_nonblank_and_has_source_csv(self):
        figures = sorted((self.output / "figures").glob("*.png"))
        self.assertEqual(len(figures), 11)
        for figure in figures:
            source = self.output / "data" / f"{figure.stem}.csv"
            self.assertTrue(source.is_file(), figure)
            with Image.open(figure) as image:
                self.assertGreater(image.width, 300)
                self.assertGreater(image.height, 200)
                self.assertGreater(sum(ImageStat.Stat(image.convert("RGB")).var), 1.0)

    def test_required_labels_and_exact_controls_are_present(self):
        plotting_source = (ROOT / "plotting.py").read_text(encoding="utf-8")
        self.assertIn("NON-PREDICTIVE DIMENSIONLESS SENSITIVITY ANALYSIS", plotting_source)
        caption1 = next((self.output / "captions").glob("figure01*_caption.txt")).read_text()
        self.assertIn("Conceptual workflow created by the authors; not a simulation output.", caption1)
        controls = {row["scenario_id"] for row in self._rows("figure07_controls")}
        self.assertEqual(controls, {
            "baseline_sensitivity", "zero_mixing", "spin_independent_reaction",
            "fast_relaxation", "rapid_escape", "doublet_benchmark",
            "quartet_benchmark",
        })
        ratios = {float(row["kq_over_kd"]) for row in self._rows("figure05_mixing_relaxation_heatmaps")}
        self.assertEqual(ratios, {0.0, 0.1, 1.0, 10.0})

    def test_known_publication_layout_overlaps_are_prevented(self):
        controls = plot_controls(self._rows("figure07_controls"))
        controls.canvas.draw()
        controls_axis = controls.axes[0].get_window_extent()
        controls_legend = controls.axes[0].get_legend().get_window_extent()
        self.assertGreaterEqual(controls_legend.y0, controls_axis.y1)
        plt.close(controls)

        solver = plot_solver_validation(
            self._rows("figure08_solver_validation"), TOLERANCES
        )
        solver.canvas.draw()
        solver_legend = solver.axes[1].get_legend()
        self.assertEqual(solver_legend.get_frame().get_alpha(), 1.0)
        self.assertEqual(solver_legend.get_frame().get_facecolor(), (1.0, 1.0, 1.0, 1.0))
        plt.close(solver)

        bulk = plot_bulk_rates(self._rows("figure10_bulk_rate_comparison"))
        bulk.canvas.draw()
        renderer = bulk.canvas.get_renderer()
        table = bulk.axes[1].tables[0]
        for coordinates, cell in table.get_celld().items():
            cell_box = cell.get_window_extent(renderer)
            text_box = cell.get_text().get_window_extent(renderer)
            self.assertGreaterEqual(text_box.x0, cell_box.x0, coordinates)
            self.assertLessEqual(text_box.x1, cell_box.x1, coordinates)
            self.assertGreaterEqual(text_box.y0, cell_box.y0, coordinates)
            self.assertLessEqual(text_box.y1, cell_box.y1, coordinates)
        plt.close(bulk)

    def test_sweep_rows_have_inputs_provenance_results_and_balances(self):
        required = {
            "scenario_id", "initial_state_definition", "p_doublet_initial",
            "reference_rate_s^-1", "duration_s", "mixing_over_reference",
            "local_mixing_proxy_rad_s", "radical_relaxation_s^-1",
            "oxygen_relaxation_s^-1", "k_escape_s^-1", "k_doublet_s^-1",
            "k_quartet_s^-1", "kq_over_kd", "final_doublet_population",
            "final_quartet_population", "doublet_reaction_yield_per_encounter",
            "quartet_reaction_yield_per_encounter",
            "primary_superoxide_yield_per_encounter", "escape_yield_per_encounter",
            "unresolved_probability", "probability_balance", "probability_balance_error",
            "solver", "non_predictive", "output_class", "units", "limitations",
            "source_fingerprint_sha256", "config_path", "provenance_path",
        }
        for stem in (
            "figure04_mixing_escape_heatmaps", "figure05_mixing_relaxation_heatmaps",
            "figure06_reaction_selectivity", "one_factor_at_a_time_sensitivity",
        ):
            rows = self._rows(stem)
            self.assertTrue(required.issubset(rows[0]), stem)
            for row in rows:
                self.assertEqual(row["non_predictive"], "True")
                self.assertLess(abs(float(row["probability_balance"]) - 1), 1e-8)

    def test_spin_independent_null_is_initial_state_and_mixing_independent(self):
        yields = []
        for state, p_doublet in (
            ("unpolarized", None), ("doublet", None), ("quartet", None),
            ("mixture", .25), ("mixture", .75),
        ):
            for mixing in (0.0, 10.0):
                row = encounter_result_row(
                    scenario_id="null", reference_rate_s=1e6,
                    mixing_over_reference=mixing,
                    radical_relaxation_over_reference=.3,
                    oxygen_relaxation_over_reference=.2,
                    escape_over_reference=1.0, kq_over_kd=1.0,
                    initial_state=state, p_doublet=p_doublet,
                )
                yields.append(row["primary_superoxide_yield_per_encounter"])
        self.assertLess(max(yields) - min(yields), 2e-12)

    def test_figure_baseline_data_match_independent_recomputation(self):
        row = self._rows("figure03_initial_state_comparison")[0]
        recomputed = encounter_result_row(
            scenario_id="check", reference_rate_s=float(row["reference_rate_s^-1"]),
            mixing_over_reference=float(row["mixing_over_reference"]),
            radical_relaxation_over_reference=float(row["radical_relaxation_over_reference"]),
            oxygen_relaxation_over_reference=float(row["oxygen_relaxation_over_reference"]),
            escape_over_reference=float(row["escape_over_reference"]),
            kq_over_kd=float(row["kq_over_kd"]), initial_state=row["initial_state"],
            p_doublet=None if row["p_doublet_requested"] == "" else float(row["p_doublet_requested"]),
        )
        for field in (
            "doublet_reaction_yield_per_encounter",
            "quartet_reaction_yield_per_encounter",
            "primary_superoxide_yield_per_encounter",
            "escape_yield_per_encounter", "unresolved_probability",
        ):
            self.assertAlmostEqual(float(row[field]), recomputed[field], places=12)

    def test_downstream_rows_are_finite_nonnegative_and_balanced(self):
        for row in self._rows("figure09_downstream_kinetics"):
            values = [float(row[field]) for field in (
                "radical_pool_m", "ho2_m", "o2minus_m", "hydrogen_peroxide_m",
                "accumulated_hydrogen_peroxide_loss_m",
            )]
            self.assertTrue(np.all(np.isfinite(values)))
            self.assertGreaterEqual(min(values), 0)
            self.assertLess(float(row["radical_equivalent_balance_error_m"]), 1e-12)

    def test_bulk_reaction_identity_and_exact_condition_matching(self):
        authority = validate_authority_bundle(
            ROOT / "configs" / "doxorubicin_parameters.json",
            ROOT / "configs" / "parameter_provenance.csv",
        )
        matched = evaluate_bulk_superoxide_rate(
            authority, "doxorubicin_semiquinone_plus_oxygen_pH6", 1e-6, 2e-4,
            use_scope=USE_CONDITION_MATCHED,
            condition_profile=authority["condition_profiles"]["land_1985_pH6"],
        )
        self.assertAlmostEqual(matched["rate_m_s"], .07)
        with self.assertRaises(ModelPolicyError):
            evaluate_bulk_superoxide_rate(
                authority, "cuzn_sod_superoxide", 1e-6, 2e-4,
                use_scope=USE_CONDITION_MATCHED,
                condition_profile=authority["condition_profiles"]["land_1985_pH6"],
            )

    def test_solver_and_circuit_records_are_honest(self):
        solver = self._rows("table06_independent_solver_validation", "tables")
        self.assertTrue(all(row["passes"] == "True" for row in solver))
        self.assertLess(max(float(row["worst_absolute_error"]) for row in solver), 3e-9)
        circuit = self._rows("table07_circuit_validation", "tables")
        self.assertTrue(all(row["passes"] == "True" for row in circuit))
        self.assertEqual(circuit[0]["qiskit_execution_status"], "executed")
        self.assertTrue((self.output / "figures" / "figure11_circuit_validation.png").is_file())

    def test_word_document_is_generated_from_pipeline_tables_and_figures(self):
        output = Path(self.temporary.name) / "handoff.docx"
        build_docx(self.output, output)
        self.assertTrue(output.is_file())
        self.assertGreater(output.stat().st_size, 100_000)
        document = Document(output)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn(PROVENANCE_STATEMENT, text)
        self.assertIn("Supported Claims", text)
        self.assertGreaterEqual(len(document.inline_shapes), 11)
        self.assertGreaterEqual(len(document.tables), 12)

    def test_manifest_hashes_resolve_and_overwrite_is_explicit(self):
        manifest = json.loads((self.output / "metadata" / "run_manifest.json").read_text())
        self.assertIn("command_line_invocation", manifest)
        self.assertIn("grid_definition", manifest)
        self.assertIn("numerical_tolerances", manifest)
        self.assertIn("figure_source_and_render_linkage", manifest)
        self.assertEqual(manifest["random_seeds"]["circuit_state"], 1729)
        for relative, expected in manifest["generated_file_hashes_sha256"].items():
            path = self.output / relative
            self.assertTrue(path.is_file(), relative)
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), expected)
        refused = subprocess.run(
            [sys.executable, "paper_analysis.py", "--quick", "--formats", "png",
             "--output-dir", str(self.output)],
            cwd=ROOT, text=True, capture_output=True,
        )
        self.assertNotEqual(refused.returncode, 0)
        self.assertIn("--overwrite", refused.stderr)


if __name__ == "__main__":
    unittest.main()
